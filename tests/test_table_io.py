import pytest
from openpyxl import load_workbook

from wiresheet.colors import parse_color
from wiresheet.components import EndKind, parse_kind
from wiresheet.demo import demo_cables
from wiresheet.model import End
from wiresheet.table_io import (HEADERS, SHEET_NAME, example_cables, import_table,
                               rows_to_cables, write_table)


def table(*rows):
    return [HEADERS, *rows]


# --- andata e ritorno -----------------------------------------------------------

@pytest.mark.parametrize("suffix", [".xlsx", ".csv"])
def test_round_trip(tmp_path, suffix):
    path = tmp_path / f"cavi{suffix}"
    write_table(path, demo_cables())
    result = import_table(path)
    assert result.errors == []
    assert result.cables == demo_cables()


@pytest.mark.parametrize("suffix", [".xlsx", ".csv"])
def test_template_contains_valid_examples(tmp_path, suffix):
    path = tmp_path / f"template{suffix}"
    write_table(path, example_cables())
    result = import_table(path)
    assert result.errors == []
    assert result.cables == example_cables()
    assert all(c.validate() == [] for c in result.cables)


def test_xlsx_template_layout(tmp_path):
    path = tmp_path / "template.xlsx"
    write_table(path, example_cables())
    wb = load_workbook(path)
    ws = wb[SHEET_NAME]
    assert [c.value for c in ws[1]] == HEADERS
    assert ws.freeze_panes == "A2"
    assert len(ws.data_validations.dataValidation) >= 4  # tendine componente/variante/pin/colore
    assert "Istruzioni" in wb.sheetnames


def test_csv_uses_semicolon_and_bom(tmp_path):
    path = tmp_path / "cavi.csv"
    write_table(path, example_cables())
    raw = path.read_bytes()
    assert raw.startswith(b"\xef\xbb\xbf")
    assert raw.decode("utf-8-sig").splitlines()[0] == ";".join(HEADERS)


def test_csv_with_comma_delimiter(tmp_path):
    path = tmp_path / "cavi.csv"
    path.write_text(",".join(HEADERS) + "\nC1,JST XH,2,1,RD,300,1,JST XH,2\n", encoding="utf-8")
    result = import_table(path)
    assert result.errors == [] and len(result.cables) == 1


# --- interpretazione dei valori -----------------------------------------------

def test_continuation_rows_inherit_ends_and_groups_by_name():
    result = rows_to_cables(table(
        ["A", "JST XH", 3, 1, "RD", 100, "", "Puntalini", ""],
        ["B", "Forcelle", "", "", "BK", 50, "", "Taglio netto", ""],
        ["A", "", "", 2, "BK", 100, "", "", ""],
    ))
    assert result.errors == []
    a, b = result.cables
    assert a.name == "A" and len(a.wires) == 2 and a.p2 == End(EndKind.FERRULE)
    assert b.p1 == End(EndKind.FORK) and b.wires[0].p1_pin is None


def test_lenient_values():
    result = rows_to_cables(table(
        ["A", "jst_xh", "4 pin", "1", "rosso", "12,5", "", "flat cut", ""],
        ["A", "", "", 2.0, "giallo-verde", 100.0, "", "", ""],
    ))
    assert result.errors == []
    cable = result.cables[0]
    assert cable.p1 == End(EndKind.JST_XH, 4)
    assert [w.color for w in cable.wires] == ["RD", "YE/GN"]
    assert cable.wires[0].length_mm == 12.5


def test_english_headers_accepted():
    rows = [["name", "p1 type", "p1 positions", "p1 pin", "color", "length", "p2 pin", "p2 type",
             "p2 positions"],
            ["A", "JST XH", 2, 1, "RD", 100, 2, "JST XH", 2]]
    result = rows_to_cables(rows)
    assert result.errors == [] and result.cables[0].wires[0].p2_pin == 2


def test_empty_rows_ignored():
    result = rows_to_cables(table([None] * 9, ["A", "Forcelle", "", "", "RD", 10, "", "Forcelle", ""],
                                  [""] * 9))
    assert result.errors == [] and len(result.cables) == 1


# --- errori con numero di riga ----------------------------------------------------

def test_missing_columns():
    result = rows_to_cables([["Cavo", "Colore"], ["A", "RD"]])
    assert result.cables == []
    assert "Colonne mancanti" in result.errors[0] and "P1 componente" in result.errors[0]


def test_errors_report_row_numbers():
    result = rows_to_cables(table(
        ["", "JST XH", 2, 1, "RD", 100, 1, "JST XH", 2],          # riga 2: nome mancante
        ["A", "Boh", "", 1, "RD", 100, "", "Forcelle", ""],          # riga 3: componente
        ["B", "JST XH", 9, 1, "RD", 100, "", "Forcelle", ""],        # riga 4: pin non supportati
        ["C", "JST XH", 2, 1, "XX", -5, "", "Forcelle", ""],         # riga 5: colore + lunghezza
        ["D", "JST XH", 2, "", "RD", 100, "", "Forcelle", ""],       # riga 6: pin P1 mancante
        ["E", "JST XH", 2, 1, "RD", 100, "", "Forcelle", ""],
        ["E", "Forcelle", "", "", "BK", 100, "", "", ""],            # riga 8: P1 diverso
    ))
    text = "\n".join(result.errors)
    for expected in ("Riga 2: nome", "Riga 3, P1: componente 'Boh'", "Riga 4, P1: JST XH a 9",
                     "Riga 5: colore 'XX'", "Riga 5: lunghezza", "Riga 6: pin P1",
                     "Riga 8: P1 diverso"):
        assert expected in text, expected


def test_first_row_without_component():
    result = rows_to_cables(table(["A", "", "", "", "RD", 100, "", "Forcelle", ""]))
    assert "componente P1 mancante" in result.errors[0]


def test_empty_file():
    assert rows_to_cables([]).errors == ["Il file è vuoto."]


# --- parser di colori e componenti -------------------------------------------------

@pytest.mark.parametrize("text,code", [("RD", "RD"), ("rd", "RD"), ("Rosso", "RD"),
                                       ("YE/GN", "YE/GN"), ("GN/YE", "YE/GN"),
                                       ("giallo/verde", "YE/GN"), ("Nero", "BK"), ("??", None)])
def test_parse_color(text, code):
    assert parse_color(text) == code


@pytest.mark.parametrize("text,kind", [("JST XH", EndKind.JST_XH), ("xh", EndKind.JST_XH),
                                       ("Puntalino", EndKind.FERRULE), ("fork", EndKind.FORK),
                                       ("Taglio netto", EndKind.FLAT_CUT), ("boh", None)])
def test_parse_kind(text, kind):
    assert parse_kind(text) is kind


def test_import_gm_variants():
    result = rows_to_cables(table(
        ["M", "GM-2112", "s5a", "", "RD", 300, 1, "GM1312", "P2"],
        ["M", "", "", "", "BK", 300, 2, "", ""],
    ))
    assert result.errors == [
        "Riga 2: pin P1 mancante o non valido ('').",
        "Riga 3: pin P1 mancante o non valido ('').",
    ]
    ok = rows_to_cables(table(["M", "GM2112", "S5A", 1, "RD", 300, 1, "GM1312", "P2"]))
    assert ok.errors == []
    assert ok.cables[0].p1 == End(EndKind.GM2112, 5, "S5A")
    assert ok.cables[0].p2 == End(EndKind.GM1312, 2, "P2")


@pytest.mark.parametrize("variant, message", [
    ("", "variante mancante"), ("S5", "non valida"), (4, "non valida"),
])
def test_import_gm_bad_variant(variant, message):
    result = rows_to_cables(table(["M", "GM2112", variant, 1, "RD", 300, "", "Puntalini", ""]))
    assert result.cables == [] or result.errors
    assert any(message in e for e in result.errors)


def test_import_terminal_block_and_faston():
    ok = rows_to_cables(table(["T", "Morsettiera 5.08", 8, 8, "RD", 300, "", "Faston", ""]))
    assert ok.errors == []
    assert ok.cables[0].p1 == End(EndKind.TERMINAL_BLOCK, 8)
    assert ok.cables[0].p2 == End(EndKind.FASTON_F)
    bad = rows_to_cables(table(["T", "Morsettiera", 9, 1, "RD", 300, "", "Faston", ""]))
    assert any("Morsettiera 5.08 a 9 pin non supportato (2-8)" in e for e in bad.errors)


# --- cavi a piu' rami -------------------------------------------------------------

def test_branches_by_name_and_component():
    result = rows_to_cables(table(
        ["P", "GM1712", "S7", 2, "BU", 500, 2, "JST XH", 3, "Endstop X"],
        ["P", "", "", 3, "RD", 500, 1, "JST XH", 3, "Servo probe"],
        ["P", "", "", 1, "WH", 500, 3, "", "", "Endstop X"],      # continua Endstop X
        ["P", "", "", 4, "BK", 500, 2, "", "", "Servo probe"],
    ))
    assert result.errors == []
    cable = result.cables[0]
    assert [(b.name, len(b.wires)) for b in cable.branches] == [("Endstop X", 2), ("Servo probe", 2)]
    assert cable.validate() == []


def test_unnamed_branches_split_by_component():
    # come AC IN: due puntalini e un filo a taglio netto, senza nome del ramo
    result = rows_to_cables(table(
        ["AC IN", "Faston", "", "", "BU", 200, "", "Puntalini", ""],
        ["AC IN", "", "", "", "BN", 200, "", "", ""],
        ["AC IN", "", "", "", "YE", 500, "", "Taglio netto", ""],
    ))
    assert result.errors == []
    cable = result.cables[0]
    assert [(b.end.kind, len(b.wires)) for b in cable.branches] == [
        (EndKind.FERRULE, 2), (EndKind.FLAT_CUT, 1)]


def test_new_branch_without_component():
    result = rows_to_cables(table(
        ["P", "JST XH", 2, 1, "RD", 100, "", "Puntalini", "", "A"],
        ["P", "", "", 2, "BK", 100, "", "", "", "B"],
    ))
    assert result.errors == ["Riga 3: componente P2 mancante (va indicato sulla prima riga del ramo 'B')."]


def test_branch_column_written_on_every_row():
    from wiresheet.table_io import cables_to_rows
    fans = next(c for c in demo_cables() if c.name == "FANS, MAT SENSOR")
    rows = cables_to_rows([fans])
    assert [r[9] for r in rows] == ["Cooling fan"] * 2 + ["Material sensor"] * 3 + ["Extruder fan"] * 2
    assert [r[7] for r in rows] == ["JST XH", "", "Morsettiera 5.08", "", "", "Taglio netto", ""]


# --- sezione del filo ---------------------------------------------------------------

def test_section_in_mm2_awg_or_both():
    result = rows_to_cables(table(
        ["S", "Forcelle", "", "", "RD", 100, "", "Puntalini", "", "", "0,25"],
        ["S", "", "", "", "BK", 100, "", "", "", "", "22AWG"],
        ["S", "", "", "", "WH", 100, "", "", "", "", "1.5mm² (15AWG)"],
        ["S", "", "", "", "BU", 100, "", "", "", "", "AWG 20"],
        ["S", "", "", "", "GN", 100, "", "", "", "", ""],
    ))
    assert result.errors == []
    assert [(w.section_mm2, w.awg) for w in result.cables[0].wires] == [
        (0.25, None), (None, 22), (1.5, 15), (None, 20), (None, None)]


def test_bad_section():
    result = rows_to_cables(table(
        ["S", "Forcelle", "", "", "RD", 100, "", "Puntalini", "", "", "tanto"],
        ["S", "", "", "", "BK", 100, "", "", "", "", "99AWG"]))
    assert result.errors == [
        "Riga 2: sezione 'tanto' non valida (es. 0.25mm² (24AWG), 0,25 oppure 24AWG).",
        "Riga 3: sezione '99AWG' non valida (es. 0.25mm² (24AWG), 0,25 oppure 24AWG)."]


def test_section_round_trip(tmp_path):
    from wiresheet.model import Cable, Wire
    cables = [Cable("S", End(EndKind.FORK), End(EndKind.FERRULE), [
        Wire("RD", 100, None, None, 0.25, None), Wire("BK", 100, None, None, None, 22),
        Wire("WH", 100, None, None, 1.5, 15), Wire("GN", 100)])]
    for suffix in (".csv", ".xlsx"):
        write_table(tmp_path / f"s{suffix}", cables)
        result = import_table(tmp_path / f"s{suffix}")
        assert result.errors == [] and result.cables == cables
