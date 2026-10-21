"""Import/export dei cavi in tabella (Excel .xlsx o CSV): una riga per filo.

Le righe con lo stesso valore in "Cavo" formano un cavo. P1 (componente e variante)
si compila solo sulla prima riga del cavo: le righe successive lo ereditano (se
compilato, deve coincidere).

Cavi a piu' rami: un ramo e' identificato da "Ramo" (nome, anche vuoto) + componente P2.
P2 si compila sulla prima riga di ogni ramo; le righe con P2 vuoto continuano l'ultimo
ramo con lo stesso nome.
"""

import csv
import re
from dataclasses import dataclass
from pathlib import Path

from wiresheet.colors import COLORS, parse_color
from wiresheet.gauges import EQUIVALENTS, parse_section, section_cell, section_text
from wiresheet.components import (GM_LAYOUTS, KIND_LABELS, POSITION_RANGES, EndKind,
                                  gm_variants, parse_gm_variant, parse_kind,
                                  variant_positions)
from wiresheet.model import Branch, Cable, End, Wire, end_variants

SHEET_NAME = "Cavi"
CSV_DELIMITER = ";"  # quello che Excel usa con le impostazioni italiane

# (chiave interna, intestazione nel file, alias accettati in import)
COLUMNS = [
    ("cable", "Cavo", ("cable", "nome", "name")),
    ("p1_kind", "P1 componente", ("p1", "p1 tipo", "p1 type", "p1 kind")),
    ("p1_variant", "P1 variante", ("p1 pin totali", "p1 positions", "p1 variant")),
    ("p1_pin", "Pin P1", ("p1 pin",)),
    ("color", "Colore", ("color", "colour")),
    ("length", "Lunghezza (mm)", ("lunghezza", "length", "length (mm)", "length mm")),
    ("p2_pin", "Pin P2", ("p2 pin",)),
    ("p2_kind", "P2 componente", ("p2", "p2 tipo", "p2 type", "p2 kind")),
    ("p2_variant", "P2 variante", ("p2 pin totali", "p2 positions", "p2 variant")),
    ("branch", "Ramo", ("branch", "funzione", "funzione/segnale")),
    ("section", "Sezione", ("sezione (mm²)", "sezione (mm2)", "sezione mm2", "section", "mm2", "mm²",
                            "awg", "gauge")),
]
HEADERS = [header for _, header, _ in COLUMNS]
REQUIRED = {"cable", "p1_kind", "color", "length", "p2_kind"}


def example_cables() -> list[Cable]:
    """Cavi di esempio inseriti nel template."""
    xh = lambda n: End(EndKind.JST_XH, n)
    return [
        Cable("CABLE1", xh(2), xh(2), [Wire("RD", 350, 1, 1), Wire("BK", 350, 2, 2)]),
        Cable("CABLE2", xh(4), End(EndKind.FERRULE),
              [Wire("RD", 200, 1), Wire("BK", 200, 2), Wire("YE", 200, 3), Wire("GN", 200, 4)]),
    ]


# --- export ---------------------------------------------------------------------

def cables_to_rows(cables: list[Cable]) -> list[list]:
    """Righe della tabella (senza intestazione). P1 solo sulla prima riga del cavo,
    P2 sulla prima riga di ogni ramo, il nome del ramo su tutte le sue righe."""
    rows = []
    for cable in cables:
        first_cable = True
        for branch in cable.branches:
            end = branch.end
            for i, w in enumerate(branch.wires):
                first = i == 0
                rows.append([
                    cable.name,
                    KIND_LABELS[cable.p1.kind] if first_cable else "",
                    _variant(cable.p1) if first_cable else "",
                    w.p1_pin if cable.p1.is_connector else "",
                    w.color,
                    _number(w.length_mm),
                    w.p2_pin if end.is_connector else "",
                    KIND_LABELS[end.kind] if first else "",
                    _variant(end) if first else "",
                    branch.name,
                    section_cell(w.section_mm2, w.awg),
                ])
                first_cable = False
    return rows


def _variant(end: End):
    if end.variant:
        return end.variant
    return end.positions if end.is_connector else ""


def _number(value: float):
    return int(value) if float(value).is_integer() else value


def write_table(path: str | Path, cables: list[Cable]) -> None:
    """Scrive .xlsx o .csv a seconda dell'estensione."""
    path = Path(path)
    if path.suffix.lower() == ".csv":
        write_csv(path, cables)
    else:
        write_xlsx(path, cables)


def write_csv(path: str | Path, cables: list[Cable]) -> None:
    # utf-8 con BOM: Excel lo apre con gli accenti corretti
    with open(path, "w", newline="", encoding="utf-8-sig") as f:
        writer = csv.writer(f, delimiter=CSV_DELIMITER)
        writer.writerow(HEADERS)
        writer.writerows(cables_to_rows(cables))


def write_xlsx(path: str | Path, cables: list[Cable]) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation

    wb = Workbook()
    ws = wb.active
    ws.title = SHEET_NAME
    ws.append(HEADERS)
    for row in cables_to_rows(cables):
        ws.append(row)

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="2F6FDF")
    for cell in ws[1]:
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.freeze_panes = "A2"
    for col, width in zip("ABCDEFGHIJK", (14, 16, 12, 9, 10, 15, 9, 16, 12, 22, 17)):
        ws.column_dimensions[col].width = width

    last = 1000
    kinds = ",".join(KIND_LABELS.values())
    all_variants = [e.variant or str(e.positions) for k in EndKind for e in end_variants(k)
                    if e.is_connector]
    variants = ",".join(dict.fromkeys(all_variants))
    max_pins = max(e.positions for k in EndKind for e in end_variants(k))
    pins = ",".join(str(n) for n in range(1, max_pins + 1))
    colors = ",".join(COLORS)
    for columns, values in (("BH", kinds), ("CI", variants), ("DG", pins), ("E", colors)):
        dv = DataValidation(type="list", formula1=f'"{values}"', allow_blank=True)
        dv.error = "Valore non valido: scegli dalla lista."
        dv.errorTitle = "WireSheet"
        ws.add_data_validation(dv)
        for col in columns:
            dv.add(f"{col}2:{col}{last}")
    length_dv = DataValidation(type="decimal", operator="greaterThan", formula1="0",
                               allow_blank=True)
    length_dv.error = "La lunghezza deve essere un numero maggiore di 0."
    ws.add_data_validation(length_dv)
    length_dv.add(f"F2:F{last}")
    # sezione: tendina con le misure comuni, ma si puo' scrivere anche un valore diverso
    sections = ",".join(section_text(mm2, awg) for mm2, awg in EQUIVALENTS)
    section_dv = DataValidation(type="list", formula1=f'"{sections}"', allow_blank=True,
                                showErrorMessage=False)
    ws.add_data_validation(section_dv)
    section_dv.add(f"K2:K{last}")

    _write_instructions(wb.create_sheet("Istruzioni"))
    wb.save(path)


def _write_instructions(ws) -> None:
    from openpyxl.styles import Font

    lines = [
        ("Come compilare il foglio \"Cavi\"", True),
        ("", False),
        ("Una riga per ogni filo. Le righe con lo stesso valore in \"Cavo\" formano un cavo.", False),
        ("Componente e variante di P1 bastano sulla prima riga del cavo.", False),
        ("Cavi a più rami (un ingresso P1, più uscite): le righe con lo stesso \"Ramo\" e lo stesso", False),
        ("    componente P2 formano un ramo. P2 basta sulla prima riga del ramo.", False),
        ("", False),
        ("Colonne", True),
        ("Cavo: nome del cavo, es. CABLE1 (obbligatorio).", False),
        (f"P1 / P2 componente: {', '.join(KIND_LABELS.values())}.", False),
        ("P1 / P2 variante: per JST XH il numero di pin (2-7), per Morsettiera 5.08 (2-8).", False),
        ("    Per GM1312/GM1712/GM2112 (a pannello) genere + pin: S = femmina, P = maschio, es. S4, P7, S5A.", False),
        *[(f"    {KIND_LABELS[k]}: " + ", ".join(gm_variants(k)), False) for k in GM_LAYOUTS],
        ("Pin P1 / Pin P2: pin del connettore a cui va il filo. Vuoto per puntalini, forcelle, faston, taglio netto.", False),
        ("Colore: " + ", ".join(f"{c.code} ({c.name_it})" for c in COLORS.values()) + ".", False),
        ("Lunghezza (mm): lunghezza del filo in millimetri.", False),
        ("Ramo: nome del ramo, es. Cooling fan (facoltativo; vuoto per i cavi semplici).", False),
        ("Sezione: sezione del filo (facoltativa), es. 0.25mm² (24AWG). Basta una delle due unità (0,25 oppure", False),
        ("    24AWG): l'altra si ricava dalla tabella di equivalenza (0,25 = 24AWG, 0,5 = 20AWG, 1,5 = 16AWG…).", False),
        ("", False),
        ("Il file si importa da WireSheet con File → Importa cavi da Excel/CSV.", False),
    ]
    for text, bold in lines:
        ws.append([text])
        if bold:
            ws.cell(ws.max_row, 1).font = Font(bold=True, size=12)
    ws.column_dimensions["A"].width = 110


# --- import ---------------------------------------------------------------------

@dataclass
class ImportResult:
    cables: list[Cable]
    errors: list[str]


def read_rows(path: str | Path) -> list[list]:
    """Tutte le righe del file (intestazione compresa) come liste di valori."""
    path = Path(path)
    if path.suffix.lower() == ".csv":
        text = path.read_text(encoding="utf-8-sig")
        try:
            dialect = csv.Sniffer().sniff(text.splitlines()[0] if text else "", delimiters=";,\t")
            delimiter = dialect.delimiter
        except csv.Error:
            delimiter = CSV_DELIMITER
        return [row for row in csv.reader(text.splitlines(), delimiter=delimiter)]

    from openpyxl import load_workbook
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb[SHEET_NAME] if SHEET_NAME in wb.sheetnames else wb.worksheets[0]
    rows = [list(row) for row in ws.iter_rows(values_only=True)]
    wb.close()
    return rows


def import_table(path: str | Path) -> ImportResult:
    return rows_to_cables(read_rows(path))


def _norm(text) -> str:
    return " ".join(str(text).strip().lower().split()) if text is not None else ""


def _column_map(header: list) -> dict[str, int]:
    lookup = {}
    for key, title, aliases in COLUMNS:
        for name in (title, key, *aliases):
            lookup[_norm(name)] = key
    found = {}
    for index, cell in enumerate(header):
        key = lookup.get(_norm(cell))
        if key and key not in found:
            found[key] = index
    return found


def rows_to_cables(rows: list[list]) -> ImportResult:
    rows = [r for r in rows if any(_norm(c) for c in r)]
    if not rows:
        return ImportResult([], ["Il file è vuoto."])
    columns = _column_map(rows[0])
    missing = [title for key, title, _ in COLUMNS if key in REQUIRED and key not in columns]
    if missing:
        return ImportResult([], [f"Colonne mancanti: {', '.join(missing)}."])

    errors: list[str] = []
    cables: dict[str, Cable] = {}
    for line, row in enumerate(rows[1:], start=2):
        get = lambda key: row[columns[key]] if key in columns and columns[key] < len(row) else None
        name = str(get("cable") or "").strip()
        if not name:
            errors.append(f"Riga {line}: nome del cavo mancante.")
            continue

        ends = []
        for side in ("p1", "p2"):
            end, error = _parse_end(get(f"{side}_kind"), get(f"{side}_variant"))
            if error:
                errors.append(f"Riga {line}, {side.upper()}: {error}")
            ends.append(end)

        branch_name = str(get("branch") or "").strip()
        cable = cables.get(name)
        if cable is None:
            if None in ends:
                missing_side = "P1" if ends[0] is None else "P2"
                if not any(f"Riga {line}, {missing_side}" in e for e in errors):
                    errors.append(f"Riga {line}: componente {missing_side} mancante "
                                  f"(va indicato sulla prima riga del cavo {name}).")
                continue
            cable = cables[name] = Cable(name, ends[0], branches=[])
        elif ends[0] is not None and ends[0] != cable.p1:
            errors.append(f"Riga {line}: P1 diverso dalla prima riga del cavo {name}.")

        branch = _find_branch(cable, branch_name, ends[1])
        if branch is None:
            if ends[1] is None:
                if not any(f"Riga {line}, P2" in e for e in errors):
                    what = f"del ramo '{branch_name}'" if branch_name else f"del cavo {name}"
                    errors.append(f"Riga {line}: componente P2 mancante "
                                  f"(va indicato sulla prima riga {what}).")
                continue
            branch = Branch(ends[1], name=branch_name)
            cable.branches.append(branch)

        wire, wire_errors = _parse_wire(cable.p1, branch.end, get)
        errors += [f"Riga {line}: {e}" for e in wire_errors]
        if wire is not None:
            branch.wires.append(wire)

    cables = {k: c for k, c in cables.items() if c.branches}
    return ImportResult(list(cables.values()), errors)


def _find_branch(cable: Cable, name: str, end: End | None) -> Branch | None:
    """Ramo con questo nome e questa estremita' (P2 vuoto: l'ultimo con questo nome)."""
    for branch in reversed(cable.branches):
        if branch.name == name and (end is None or branch.end == end):
            return branch
    return None


def _parse_end(kind_value, variant_value) -> tuple[End | None, str | None]:
    if not _norm(kind_value):
        return None, None  # riga di continuazione: eredita dalla prima riga
    kind = parse_kind(str(kind_value))
    if kind is None:
        return None, (f"componente '{kind_value}' sconosciuto "
                      f"(valori ammessi: {', '.join(KIND_LABELS.values())}).")
    if kind in GM_LAYOUTS:
        if not _norm(variant_value):
            return None, f"variante mancante per {KIND_LABELS[kind]} (es. S4 femmina, P4 maschio)."
        variant = parse_gm_variant(kind, variant_value)
        if variant is None:
            return None, (f"variante '{variant_value}' non valida per {KIND_LABELS[kind]} "
                          f"(valori ammessi: {', '.join(gm_variants(kind))}).")
        return End(kind, variant_positions(variant), variant), None
    if kind not in POSITION_RANGES:
        return End(kind), None
    allowed = POSITION_RANGES[kind]
    positions = _parse_int(variant_value)
    if positions is None:
        return None, f"variante (numero di pin) mancante per {KIND_LABELS[kind]}."
    if positions not in allowed:
        return None, (f"{KIND_LABELS[kind]} a {positions} pin non supportato "
                      f"({allowed[0]}-{allowed[-1]}).")
    return End(kind, positions), None


def _parse_wire(p1: End, p2: End, get) -> tuple[Wire | None, list[str]]:
    errors = []
    color_value = get("color")
    color = parse_color(str(color_value)) if _norm(color_value) else None
    if color is None:
        errors.append(f"colore '{color_value or ''}' sconosciuto." if _norm(color_value)
                      else "colore mancante.")
    length = _parse_float(get("length"))
    if length is None or length <= 0:
        errors.append(f"lunghezza '{get('length') or ''}' non valida.")
    section = parse_section(get("section"))
    if section is None:
        errors.append(f"sezione '{get('section')}' non valida (es. 0.25mm² (24AWG), 0,25 oppure 24AWG).")
        section = (None, None)
    pins = []
    for side, end in (("p1", p1), ("p2", p2)):
        value = get(f"{side}_pin")
        pin = _parse_int(value) if end.is_connector else None
        if end.is_connector and pin is None:
            errors.append(f"pin {side.upper()} mancante o non valido ('{value or ''}').")
        pins.append(pin)
    if errors:
        return None, errors
    return Wire(color, length, pins[0], pins[1], *section), []


def _parse_int(value) -> int | None:
    if value is None:
        return None
    if isinstance(value, (int, float)) and float(value).is_integer():
        return int(value)
    match = re.search(r"\d+", str(value))  # accetta anche "4 pin"
    return int(match.group()) if match else None


def _parse_float(value) -> float | None:
    if value is None:
        return None
    if isinstance(value, (int, float)):
        return float(value)
    try:
        return float(str(value).strip().replace(",", "."))  # accetta "12,5"
    except ValueError:
        return None
